"""
Vérification rapide de l'OCR manuscrit (prénom / nom) sur FORM1.

Compare les noms lus par le pipeline aux vérités terrain (xlsx fournis),
sans lancer tout main.py. Affiche, pour chaque copie, prédiction vs vérité,
et un score global (exact + par caractère).

Usage (depuis la racine du projet) :
    python check_names.py
    python check_names.py photo      # teste les photos au lieu des PDFs
    python check_names.py pdf 10     # limite à 10 copies
"""

import os
import re
import sys

import cv2
import numpy as np
import openpyxl

# Console Windows souvent en cp1252 : forcer UTF-8 si possible (Python >= 3.7).
try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from utils.image_io import imread_robust
from utils.template_register import get_photo_template
from utils.grid_decoder import normalize_page, set_photo_template
from utils.page1_parser import read_firstname, read_name

try:
    import fitz
except ImportError:
    fitz = None

DATA = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                    "PROJECT 2026 -DATABASE-20260518", "FORM1")


def render_pdf_p1(path, dpi=150):
    doc = fitz.open(path)
    pix = doc[0].get_pixmap(dpi=dpi)
    arr = np.frombuffer(pix.samples, np.uint8).reshape(pix.h, pix.w, pix.n)
    arr = cv2.cvtColor(arr, cv2.COLOR_RGB2BGR if pix.n == 3 else cv2.COLOR_RGBA2BGR)
    doc.close()
    return arr


def truth_names(sid):
    p = os.path.join(DATA, f"EXAM_FORM1_{sid}.xlsx")
    if not os.path.isfile(p):
        return None, None
    ws = openpyxl.load_workbook(p, data_only=True)["PAGE-01"]
    return (str(ws.cell(13, 2).value or ""), str(ws.cell(14, 2).value or ""))


def char_acc(pred, truth):
    """Similarité caractère par caractère (0..1), insensible à la casse."""
    pred = re.sub(r"[^A-Za-z]", "", pred).upper()
    truth = re.sub(r"[^A-Za-z]", "", truth).upper()
    if not truth:
        return 1.0 if not pred else 0.0
    n = max(len(pred), len(truth))
    same = sum(1 for a, b in zip(pred, truth) if a == b)
    return same / n


def main():
    kind = sys.argv[1] if len(sys.argv) > 1 else "pdf"
    limit = int(sys.argv[2]) if len(sys.argv) > 2 else 9999

    set_photo_template(get_photo_template(DATA))
    if kind == "pdf" and fitz is None:
        print("pymupdf manquant pour le mode pdf.")
        return

    if kind == "photo":
        pat = re.compile(r"EXAM_FORM1_(\d+)\.(jpg|jpeg|png|JPG|JPEG)$")
    else:
        pat = re.compile(r"EXAM_FORM1_(\d+)\.pdf$")

    files = sorted(f for f in os.listdir(DATA) if pat.match(f))
    exact_fn = exact_nm = cacc_fn = cacc_nm = n = 0

    print(f"{'ID':<8}{'prenom (verite -> predit)':<36}{'nom (verite -> predit)'}")
    print("-" * 90)
    for f in files[:limit]:
        sid = pat.match(f).group(1)
        if sid == "00000":
            continue
        tf, tn = truth_names(sid)
        if tf is None:
            continue
        if kind == "photo":
            img = imread_robust(os.path.join(DATA, f))
            if img is None:
                continue
            norm = normalize_page(img, is_photo=True)
        else:
            norm = normalize_page(render_pdf_p1(os.path.join(DATA, f)),
                                  is_photo=False, use_template=True)
        pf = read_firstname(norm)
        pn = read_name(norm)
        n += 1
        ef = re.sub(r"[^A-Za-z]", "", pf).upper() == re.sub(r"[^A-Za-z]", "", tf).upper()
        en = re.sub(r"[^A-Za-z]", "", pn).upper() == re.sub(r"[^A-Za-z]", "", tn).upper()
        exact_fn += ef
        exact_nm += en
        cacc_fn += char_acc(pf, tf)
        cacc_nm += char_acc(pn, tn)
        print(f"{sid:<8}{tf+' -> '+pf:<36}{tn+' -> '+pn}")

    if n:
        print("-" * 90)
        print(f"n={n}")
        print(f"Prénom : exact {exact_fn}/{n}={100*exact_fn/n:.0f}%  | "
              f"caractères {100*cacc_fn/n:.0f}%")
        print(f"Nom    : exact {exact_nm}/{n}={100*exact_nm/n:.0f}%  | "
              f"caractères {100*cacc_nm/n:.0f}%")


if __name__ == "__main__":
    main()
