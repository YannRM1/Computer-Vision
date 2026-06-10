"""
Évaluation rapide de l'axe MANUSCRIT uniquement (§6.3) :
  - Prénom / Nom (lignes 13-14 de PAGE-01)
  - MANTISSE / EXPOSANT / UNITE (colonnes 10-12 de l'onglet EXAM)

Recalcule ces champs directement depuis les PDF (sans le pipeline complet :
pas de signature, pas d'OCR des champs imprimés) et compare aux xlsx vérité.

Usage :
  python tools/eval_manuscrit.py [FORM1|FORM2|FORM3] [--ids 19283,36912]
"""

import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl

from utils.pdf_utils import pdf_to_images
from utils.grid_decoder import normalize_page
from utils.page1_parser import read_firstname, read_name
from utils.exam_parser import parse_exam_pages, questions_to_exam_rows
from utils.config import EXAM_START_PAGE

DATA_ROOT = "PROJECT 2026 -DATABASE-20260518"


def norm_val(v):
    if v is None:
        return None
    if isinstance(v, str):
        s = v.strip()
        try:
            return float(s.replace(",", "."))
        except ValueError:
            return s.lower()
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def equal(a, b):
    na, nb = norm_val(a), norm_val(b)
    if na is None and nb is None:
        return True
    if na is None or nb is None:
        return False
    if isinstance(na, (int, float)) and isinstance(nb, (int, float)):
        return abs(na - nb) < 1e-3
    return na == nb


def main():
    form = sys.argv[1] if len(sys.argv) > 1 else "FORM1"
    only_ids = None
    if "--ids" in sys.argv:
        only_ids = set(sys.argv[sys.argv.index("--ids") + 1].split(","))

    form_dir = os.path.join(DATA_ROOT, form)
    pdfs = sorted(f for f in os.listdir(form_dir) if f.endswith(".pdf"))

    ok_names = tot_names = ok_num = tot_num = 0
    errors = []

    for pdf in pdfs:
        sid = re.match(rf"EXAM_{form}_(\d+)\.pdf", pdf).group(1)
        if only_ids and sid not in only_ids:
            continue
        truth_path = os.path.join(form_dir, f"EXAM_{form}_{sid}.xlsx")
        if not os.path.isfile(truth_path):
            continue
        twb = openpyxl.load_workbook(truth_path, data_only=True)

        imgs = pdf_to_images(os.path.join(form_dir, pdf))
        norm = normalize_page(imgs[0], use_template=True)

        # --- Prénom / Nom ------------------------------------------------
        preds = {"Prenom": read_firstname(norm), "Nom": read_name(norm)}
        tws = twb["PAGE-01"]
        for row, label in [(13, "Prenom"), (14, "Nom")]:
            t, p = tws.cell(row, 2).value, preds[label]
            ok = equal(t, p if p else None)
            tot_names += 1
            ok_names += int(ok)
            if not ok:
                errors.append((sid, label, t, p))

        # --- EXAM : mantisse / exposant / unité ---------------------------
        questions = parse_exam_pages(imgs, exam_start_page=EXAM_START_PAGE)
        rows = questions_to_exam_rows(questions)
        ews = twb["EXAM"]
        for r in range(2, ews.max_row + 1):
            q = ews.cell(r, 1).value
            if q is None:
                continue
            prod_row = rows[r - 2] if r - 2 < len(rows) else {}
            for col, key in [(10, "MANTISSE"), (11, "EXPOSANT"), (12, "UNITE")]:
                t = ews.cell(r, col).value
                p = prod_row.get(key)
                if t is None and p is None:
                    continue
                ok = equal(t, p)
                tot_num += 1
                ok_num += int(ok)
                if not ok:
                    errors.append((sid, f"Q{q}-{key}", t, p))

    print(f"\n=== Axe manuscrit — {form} ===")
    print(f"Noms (Prenom/Nom) : {ok_names}/{tot_names}  "
          f"{100*ok_names/max(1,tot_names):.1f}%")
    print(f"EXAM (M/E/U)      : {ok_num}/{tot_num}  "
          f"{100*ok_num/max(1,tot_num):.1f}%")
    tot = tot_names + tot_num
    print(f"TOTAL             : {ok_names+ok_num}/{tot}  "
          f"{100*(ok_names+ok_num)/max(1,tot):.1f}%")
    print("\n--- Erreurs ---")
    for sid, field, t, p in errors:
        print(f"{sid} {field:14s} truth={t!r} prod={p!r}")


if __name__ == "__main__":
    main()
