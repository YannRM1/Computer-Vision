"""
Balayage du seuil de détection des cases MCQ (CHECKED_INK_THRESHOLD).

Objectif : trouver, en UNE seule passe, la valeur de seuil qui maximise la
lecture des CHOIX (cases cochées) par rapport à la vérité terrain — au lieu de
relancer main.py (15 min) pour chaque valeur.

Principe :
  1. Rendre une fois les pages d'examen de chaque PDF (cache mémoire).
  2. Pour chaque seuil candidat, repositionner exam_parser.CHECKED_INK_THRESHOLD,
     re-parser les CHOIX, et comparer aux CHOIX de la vérité terrain (onglet EXAM,
     colonnes CHOIX A-H).
  3. Afficher accuracy / faux positifs / ratés par seuil, et le meilleur.

Usage (sur ta machine, avec torch/easyocr/pymupdf installés) :
    python tools/sweep_mcq_threshold.py                 # FORM1 par défaut
    python tools/sweep_mcq_threshold.py FORM1 FORM2     # plusieurs formulaires
    python tools/sweep_mcq_threshold.py FORM1 --dpi 150

Quand tu as le meilleur seuil, mets-le dans utils/config.py
(CHECKED_INK_THRESHOLD) puis relance main.py une seule fois.
"""

import os
import re
import sys

# Racine du projet (ce script est dans tools/)
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

import openpyxl

from utils.pdf_utils import pdf_to_images
from utils import exam_parser
from utils.exam_parser import parse_exam_pages

DATA_ROOT = "PROJECT 2026 -DATABASE-20260518"
LETTERS = "ABCDEFGH"
EXAM_START_PAGE = 4
THRESHOLDS = [0.08, 0.10, 0.12, 0.14, 0.16, 0.18, 0.20, 0.22, 0.25, 0.30]


def truth_choix_rows(ws):
    """Liste, en ordre, des ensembles de lettres cochées par question (vérité).
    Onglet EXAM : col 1 = QUESTION, col 2-9 = CHOIX A-H."""
    rows = []
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 1).value is None:
            continue
        checked = set()
        for i, letter in enumerate(LETTERS):
            v = ws.cell(r, 2 + i).value
            if v not in (None, 0, "0", ""):
                checked.add(letter)
        rows.append(checked)
    return rows


def load_forms(form_names, dpi):
    """Rend une fois les images d'examen + charge les CHOIX vérité de chaque PDF."""
    forms = []
    for form in form_names:
        fdir = os.path.join(DATA_ROOT, form)
        if not os.path.isdir(fdir):
            print(f"[skip] {fdir} absent")
            continue
        for f in sorted(os.listdir(fdir)):
            m = re.match(rf"EXAM_{form}_(\d+)\.pdf$", f)
            if not m:
                continue
            tx = os.path.join(fdir, f[:-4] + ".xlsx")
            if not os.path.isfile(tx):
                continue
            try:
                tw = openpyxl.load_workbook(tx, data_only=True)
            except Exception:
                continue
            if "EXAM" not in tw.sheetnames:
                continue
            try:
                imgs = pdf_to_images(os.path.join(fdir, f), dpi=dpi)
            except Exception as e:
                print(f"[skip] rendu {f}: {e}")
                continue
            forms.append((f, truth_choix_rows(tw["EXAM"]), imgs))
    return forms


def evaluate(forms, threshold):
    """Compte OK / faux positifs / ratés des CHOIX au seuil donné."""
    exam_parser.CHECKED_INK_THRESHOLD = threshold      # repositionne le seuil
    ok = tot = fp = miss = 0
    for _name, truth_rows, imgs in forms:
        preds = parse_exam_pages(imgs, exam_start_page=EXAM_START_PAGE)
        n = min(len(truth_rows), len(preds))
        for i in range(n):
            true_set = truth_rows[i]
            pred = preds[i].get("choix", {}) or {}
            pred_set = {L for L in LETTERS if pred.get(L) == 1}
            for L in LETTERS:
                t = L in true_set
                p = L in pred_set
                if not t and not p:
                    continue                            # cellule vide des deux côtés
                tot += 1
                if t == p:
                    ok += 1
                elif p and not t:
                    fp += 1
                else:
                    miss += 1
    acc = 100.0 * ok / tot if tot else 0.0
    return acc, ok, tot, fp, miss


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    dpi = 150
    if "--dpi" in sys.argv:
        dpi = int(sys.argv[sys.argv.index("--dpi") + 1])
    form_names = args or ["FORM1"]

    print(f"[sweep] formulaires={form_names} dpi={dpi}")
    forms = load_forms(form_names, dpi)
    print(f"[sweep] {len(forms)} PDF chargés (rendu une seule fois)\n")
    if not forms:
        print("Aucun PDF/vérité trouvé.")
        return

    print(f"{'seuil':>7} | {'acc':>7} | {'OK':>5} {'TOT':>5} | {'FP':>4} {'MISS':>5}")
    print("-" * 46)
    best = None
    for T in THRESHOLDS:
        acc, ok, tot, fp, miss = evaluate(forms, T)
        flag = ""
        if best is None or acc > best[1]:
            best = (T, acc)
            flag = "  <- meilleur"
        print(f"{T:>7.2f} | {acc:>6.1f}% | {ok:>5} {tot:>5} | {fp:>4} {miss:>5}{flag}")
    print("-" * 46)
    print(f"\n>> Meilleur seuil : CHECKED_INK_THRESHOLD = {best[0]:.2f}  ({best[1]:.1f}%)")
    print("   Mets-le dans utils/config.py puis relance main.py.")


if __name__ == "__main__":
    main()
