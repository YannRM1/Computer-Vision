"""
Construction d'un jeu de chiffres manuscrits annotes a partir des formulaires.

Meme principe que build_letter_dataset.py, applique aux reponses numeriques :
les valeurs MANTISSE / EXPOSANT sont connues dans les xlsx verite terrain.
Pour chaque question numerique, on segmente la case (segmentation partagee
avec la lecture : utils.ocr_utils) ; si le nombre de composantes chiffres
correspond au nombre de chiffres de la verite, on apparie position a position.

Les noms de l'appariement par position s'appliquent aussi ici : seuls les cas
ou le compte concorde sont retenus, ce qui garantit la purete des labels.

Sortie : digit_dataset.npz (X: (N,28,28) float32, y: (N,) int 0..9,
groups: (N,) id etudiant) + digit_preview.png.

Usage : python training/build_digit_dataset.py
"""
import os
import re
import sys

import cv2
import numpy as np

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import openpyxl

from utils.pdf_utils import pdf_to_images
from utils.template_register import get_photo_template
from utils.grid_decoder import set_photo_template
from utils.exam_parser import (iter_question_blocks, numeric_answer_crops,
                               _find_mcq_checkboxes, _has_numerical_answer)
from utils.ocr_utils import _mantisse_comps, _exposant_digits, _to_gray
from utils import digit_cnn

DATA = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                    "PROJECT 2026 -DATABASE-20260518")

_PREVIEW = []


def truth_digits(value, integer=False):
    """Chaine des chiffres de la valeur verite ('3.75' -> '375'), ou None."""
    if value is None:
        return None
    s = str(value).strip()
    if integer:
        try:
            return str(abs(int(float(s))))
        except ValueError:
            return None
    s = re.sub(r"[^0-9]", "", s.rstrip("0").rstrip(".")
               if "." in s and s.endswith("0") else s)
    return s or None


def add_cells(cells, digits_str, sid, X, y, groups):
    if cells is None or digits_str is None:
        return False
    if len(cells) != len(digits_str):
        return False
    for cell, d in zip(cells, digits_str):
        arr = digit_cnn.prep_cell(cell)
        if arr is None:
            continue
        X.append(arr.astype(np.float32))
        y.append(int(d))
        groups.append(int(sid))
        if len(_PREVIEW) < 70:
            _PREVIEW.append((cell.copy(), d))
    return True


def main():
    set_photo_template(get_photo_template(DATA))
    X, y, groups = [], [], []
    n_q = 0

    for form in ("FORM1", "FORM2", "FORM3"):
        fdir = os.path.join(DATA, form)
        if not os.path.isdir(fdir):
            continue
        for f in sorted(os.listdir(fdir)):
            m = re.match(rf"EXAM_{form}_(\d+)\.pdf$", f)
            if not m or m.group(1) == "00000":
                continue
            sid = m.group(1)
            tp = os.path.join(fdir, f"EXAM_{form}_{sid}.xlsx")
            if not os.path.isfile(tp):
                continue
            try:
                tws = openpyxl.load_workbook(tp, data_only=True)["EXAM"]
            except Exception:
                continue
            try:
                imgs = pdf_to_images(os.path.join(fdir, f))
            except Exception:
                continue

            for q_idx, block in enumerate(iter_question_blocks(imgs)):
                t_mant = tws.cell(q_idx + 2, 10).value
                t_expo = tws.cell(q_idx + 2, 11).value
                if t_mant is None and t_expo is None:
                    continue
                header_cut = max(20, int(block.shape[0] * 0.22))
                content = block[header_cut:, :]
                if len(_find_mcq_checkboxes(content)) >= 2:
                    continue
                if not _has_numerical_answer(content):
                    continue
                m_img, e_img, _u = numeric_answer_crops(content)

                # mantisse : memes pretraitements que la lecture
                if t_mant is not None and m_img is not None:
                    g = _to_gray(m_img)
                    clahe = cv2.createCLAHE(clipLimit=5.0, tileGridSize=(2, 2))
                    comps = _mantisse_comps(clahe.apply(g))
                    cells = [c[4] for c in comps if not c[3]] if comps else None
                    if add_cells(cells, truth_digits(t_mant), sid, X, y, groups):
                        n_q += 1
                if t_expo is not None and e_img is not None:
                    g = _to_gray(e_img)
                    clahe = cv2.createCLAHE(clipLimit=4.0, tileGridSize=(2, 2))
                    cells, _neg = _exposant_digits(clahe.apply(g))
                    if add_cells(cells, truth_digits(t_expo, integer=True),
                                 sid, X, y, groups):
                        n_q += 1

    if not X:
        print("[digits] aucun echantillon !")
        return
    X = np.stack(X)
    y = np.array(y, dtype=np.int64)
    groups = np.array(groups, dtype=np.int64)
    out = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                       "digit_dataset.npz")
    np.savez_compressed(out, X=X, y=y, groups=groups)
    counts = {str(i): int((y == i).sum()) for i in range(10)}
    print(f"[digits] {len(X)} chiffres, {len(np.unique(groups))} etudiants, "
          f"{n_q} champs apparies")
    print("[digits] repartition:",
          " ".join(f"{k}:{v}" for k, v in counts.items() if v))
    print(f"[digits] sauvegarde : {out}")

    # apercu lisible : crops bruts + label
    TW, TH, LB, PER = 56, 72, 24, 14
    rows_img, row = [], []
    for cell, d in _PREVIEW:
        g = cell if cell.ndim == 2 else cv2.cvtColor(cell, cv2.COLOR_BGR2GRAY)
        t = cv2.cvtColor(cv2.resize(g, (TW, TH), interpolation=cv2.INTER_CUBIC),
                         cv2.COLOR_GRAY2BGR)
        cv2.rectangle(t, (0, 0), (TW - 1, TH - 1), (180, 180, 180), 1)
        band = np.full((LB, TW, 3), 255, np.uint8)
        cv2.putText(band, d, (TW // 2 - 7, LB - 6),
                    cv2.FONT_HERSHEY_SIMPLEX, 0.6, (30, 60, 140), 2)
        row.append(np.vstack([t, band]))
        if len(row) == PER:
            rows_img.append(np.hstack(row)); row = []
    if row:
        pad = [np.full((TH + LB, TW, 3), 255, np.uint8)] * (PER - len(row))
        rows_img.append(np.hstack(row + pad))
    if rows_img:
        canvas = cv2.copyMakeBorder(np.vstack(rows_img), 8, 8, 8, 8,
                                    cv2.BORDER_CONSTANT, value=(255, 255, 255))
        cv2.imwrite(os.path.join(os.path.dirname(os.path.abspath(__file__)),
                    "digit_preview.png"), canvas)
        print("[digits] apercu : digit_preview.png")


if __name__ == "__main__":
    main()
